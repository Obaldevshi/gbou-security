from types import SimpleNamespace
from datetime import datetime

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from app.schemas.user import ChangePasswordRequest
from app.schemas.guard_admin import GuardCreate
from app.schemas.school_admin import SchoolAdminCreate
from app.schemas.teacher_admin import TeacherCreate
from app.services.audit_log_service import AuditLogService
from app.services.report_service import ReportService


def test_password_policy_requires_eight_characters_and_a_letter():
    with pytest.raises(ValidationError):
        ChangePasswordRequest(
            current_password="OldPassword123!",
            new_password="12345678",
        )

    request = ChangePasswordRequest(
        current_password="OldPassword123!",
        new_password="onlyletters",
    )
    assert request.new_password == "onlyletters"


@pytest.mark.parametrize(
    ("schema", "payload"),
    [
        (
            GuardCreate,
            {"login": "guard", "full_name": "Иван Охранник", "password": "12345678"},
        ),
        (
            SchoolAdminCreate,
            {
                "school_id": 1,
                "login": "admin",
                "full_name": "Иван Администратор",
                "password": "12345678",
            },
        ),
        (
            TeacherCreate,
            {
                "login": "teacher",
                "full_name": "Иван Учитель",
                "password": "12345678",
                "building_id": 1,
                "class_ids": [1],
            },
        ),
    ],
)
def test_admin_created_accounts_require_at_least_one_letter(schema, payload):
    with pytest.raises(ValidationError):
        schema(**payload)


class FakeAuditRepository:
    def list(self, **_):
        return [
            SimpleNamespace(
                id=1,
                actor_id=2,
                actor=SimpleNamespace(
                    full_name="Администратор",
                    role=SimpleNamespace(value="school_admin"),
                ),
                school_id=3,
                school=SimpleNamespace(short_name="ГБОУ №1"),
                method="PATCH",
                path="/api/v1/school/students/5",
                status_code=200,
                ip_address="127.0.0.1",
                created_at=datetime(2026, 8, 16, 12, 0),
            )
        ], 1


def test_audit_service_maps_actor_without_sensitive_fields():
    page = AuditLogService(FakeAuditRepository()).list(
        school_id=3,
        page=1,
        page_size=25,
        search=None,
        method=None,
    )
    assert page.total == 1
    assert page.items[0].actor_name == "Администратор"
    assert "password" not in page.items[0].model_dump()


class FakeReportRepository:
    def schools(self):
        return [
            SimpleNamespace(
                short_name="ГБОУ №1",
                full_name="Государственная школа №1",
                address="Москва",
                is_active=True,
            )
        ]

    def users(self):
        return [SimpleNamespace(is_active=True)]


def test_reports_generate_valid_xlsx_and_pdf():
    service = ReportService(FakeReportRepository())
    xlsx = service.system_overview("xlsx")
    workbook = load_workbook(filename=__import__("io").BytesIO(xlsx))
    assert "Школы" in workbook.sheetnames
    assert workbook["Школы"]["A2"].value == "ГБОУ №1"

    pdf = service.system_overview("pdf")
    assert pdf.startswith(b"%PDF")
