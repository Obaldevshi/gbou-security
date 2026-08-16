from collections import Counter
from datetime import datetime
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.repositories.report_repository import ReportRepository


STATUS_LABELS = {"pending": "Ожидает выхода", "released": "Отпустил", "cancelled": "Отменена", "expired": "Просрочена"}
REASON_LABELS = {"parent_note": "По заявлению родителя", "health": "По состоянию здоровья", "other": "Своя причина"}


class ReportService:
    def __init__(self, repository: ReportRepository):
        self.repository = repository

    def system_overview(self, file_format: str) -> bytes:
        schools = self.repository.schools()
        users = self.repository.users()
        rows = [[school.short_name, school.full_name, school.address, "Активна" if school.is_active else "Отключена"] for school in schools]
        summary = [
            ["Школы", len(schools)],
            ["Активные школы", sum(item.is_active for item in schools)],
            ["Пользователи", len(users)],
            ["Активные пользователи", sum(item.is_active for item in users)],
        ]
        return self._xlsx("Школы", ["Краткое название", "Полное название", "Адрес", "Статус"], rows, summary) if file_format == "xlsx" else self._pdf("Сводный отчёт системы", ["Школа", "Полное название", "Адрес", "Статус"], rows, summary)

    def school_requests(self, school_id: int, file_format: str, date_from: datetime | None, date_to: datetime | None) -> bytes:
        requests = self.repository.requests(school_id, date_from, date_to)
        rows = [[
            item.created_at.strftime("%d.%m.%Y %H:%M"),
            item.student_full_name,
            item.class_name,
            item.teacher_full_name,
            REASON_LABELS[item.reason_type.value],
            STATUS_LABELS[item.status.value],
            item.scheduled_at.strftime("%d.%m.%Y %H:%M"),
        ] for item in requests]
        statuses = Counter(STATUS_LABELS[item.status.value] for item in requests)
        classes = Counter(item.class_name for item in requests)
        reasons = Counter(REASON_LABELS[item.reason_type.value] for item in requests)
        summary = [["Всего заявок", len(requests)]] + [[f"Статус: {key}", value] for key, value in statuses.items()] + [[f"Класс: {key}", value] for key, value in classes.items()] + [[f"Причина: {key}", value] for key, value in reasons.items()]
        headers = ["Создана", "Ученик", "Класс", "Учитель", "Причина", "Статус", "Время выхода"]
        return self._xlsx("Заявки", headers, rows, summary) if file_format == "xlsx" else self._pdf("Отчёт по заявкам школы", headers, rows, summary)

    @staticmethod
    def _xlsx(title: str, headers: list[str], rows: list[list], summary: list[list]) -> bytes:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = title
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0F766E")
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(48, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
        summary_sheet = workbook.create_sheet("Сводка")
        summary_sheet.append(["Показатель", "Значение"])
        for row in summary:
            summary_sheet.append(row)
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    @staticmethod
    def _pdf(title: str, headers: list[str], rows: list[list], summary: list[list]) -> bytes:
        font_path = Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
        if font_path.exists() and "DejaVu" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("DejaVu", str(font_path)))
        font = "DejaVu" if "DejaVu" in pdfmetrics.getRegisteredFontNames() else "Helvetica"
        output = BytesIO()
        document = SimpleDocTemplate(output, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm, topMargin=10 * mm, bottomMargin=10 * mm)
        styles = getSampleStyleSheet()
        styles["Title"].fontName = font
        story = [Paragraph(title, styles["Title"]), Spacer(1, 5 * mm)]
        summary_table = Table([["Показатель", "Значение"], *summary], repeatRows=1)
        summary_table.setStyle(TableStyle([("FONTNAME", (0, 0), (-1, -1), font), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F766E")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("GRID", (0, 0), (-1, -1), .25, colors.grey)]))
        story.extend([summary_table, Spacer(1, 7 * mm)])
        if rows:
            table = Table([headers, *rows], repeatRows=1, style=[("FONTNAME", (0, 0), (-1, -1), font), ("FONTSIZE", (0, 0), (-1, -1), 7), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F766E")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("GRID", (0, 0), (-1, -1), .2, colors.grey), ("VALIGN", (0, 0), (-1, -1), "TOP")])
            story.append(table)
        document.build(story)
        return output.getvalue()
